"""
generate-qr-codes.py

One QR code per physical spot (TV screen / sign). Each QR encodes
    <BASE_URL>/?start=<Node_Id>
so the app knows where the person is standing.

Nodes come straight from Book_4.xlsx (sheet " Node_Types", note the leading space).
By default only Type == "entry_point" nodes are generated (the TVs). To let people
re-synchronise their position along the way, also print codes for other node types
(e.g. stairs):  --types entry_point,stairs

BASE_URL must be your PUBLIC https address (e.g. https://campus-navigator.onrender.com).
It is mandatory on purpose: a local address like http://192.168.x.x only works for phones
on the same Wi-Fi, which is exactly the bug this script used to cause.

Usage:
    python scripts/generate-qr-codes.py https://campus-navigator.onrender.com
    BASE_URL=https://campus-navigator.onrender.com python scripts/generate-qr-codes.py --types entry_point,stairs

Requires: pip install "qrcode[pil]" openpyxl
"""

import argparse
import ipaddress
import os
import sys
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
import qrcode
from qrcode.constants import ERROR_CORRECT_Q

ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = ROOT / "data" / "Book_4.xlsx"
NODE_TYPES_SHEET = " Node_Types"  # leading space matches the workbook exactly
OUTPUT_DIR = ROOT / "public" / "qr-codes"


def validate_base_url(url: str, allow_local: bool) -> str:
    url = url.strip().rstrip("/")
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        sys.exit(f"BASE_URL must look like https://your-domain.com (got: {url!r})")

    host = parsed.hostname
    is_local = host in ("localhost",) or host.endswith(".local")
    try:
        ip = ipaddress.ip_address(host)
        is_local = is_local or ip.is_private or ip.is_loopback or ip.is_link_local
    except ValueError:
        pass

    if is_local and not allow_local:
        sys.exit(
            f"'{host}' is a local address: phones outside this Wi-Fi network could not open it.\n"
            "Deploy the app first and use its public https URL (see DEPLOY.md),\n"
            "or pass --allow-local if you only want to test on the same Wi-Fi."
        )
    if parsed.scheme != "https" and not is_local:
        print("WARNING: use https:// (phone motion sensors and many browsers require it).")
    return url


def load_nodes(types):
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    if NODE_TYPES_SHEET not in wb.sheetnames:
        sys.exit(f"Sheet {NODE_TYPES_SHEET!r} not found. Sheets: {wb.sheetnames}")

    wanted = {t.strip().lower() for t in types}
    selected = {}  # node_id -> type (first match wins, order of the sheet)
    for node_id, node_type in list(wb[NODE_TYPES_SHEET].iter_rows(values_only=True))[1:]:
        if node_id and node_type and str(node_type).strip().lower() in wanted:
            selected.setdefault(node_id, str(node_type).strip().lower())
    if not selected:
        sys.exit(f"No nodes of type {sorted(wanted)} found in {NODE_TYPES_SHEET!r}.")

    labels = {}
    if "Locations" in wb.sheetnames:
        for row in list(wb["Locations"].iter_rows(values_only=True))[1:]:
            name, node_id = row[1], row[5]
            if node_id and node_id not in labels and name:
                labels[node_id] = str(name).strip()
    if "Nodes" in wb.sheetnames:
        for row in list(wb["Nodes"].iter_rows(values_only=True))[1:]:
            node_id, description = row[0], row[4]
            if node_id and node_id not in labels and description:
                labels[node_id] = str(description).strip()

    return [(nid, ntype, labels.get(nid, nid)) for nid, ntype in selected.items()]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("base_url", nargs="?", default=os.environ.get("BASE_URL"), help="public URL of the deployed app")
    ap.add_argument("--types", default="entry_point", help="comma-separated node types (default: entry_point)")
    ap.add_argument("--allow-local", action="store_true", help="allow localhost / 192.168.x.x (same Wi-Fi only!)")
    args = ap.parse_args()

    if not args.base_url:
        sys.exit("Missing BASE_URL. Example:\n  python scripts/generate-qr-codes.py https://campus-navigator.onrender.com")
    base = validate_base_url(args.base_url, args.allow_local)

    nodes = load_nodes(args.types.split(","))
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    for i, (node_id, ntype, label) in enumerate(nodes, start=1):
        url = f"{base}/?start={node_id}"
        prefix = "tv" if ntype == "entry_point" else "qr"
        path = OUTPUT_DIR / f"{prefix}-{i}-{node_id}.png"
        qr = qrcode.QRCode(error_correction=ERROR_CORRECT_Q, box_size=12, border=3)
        qr.add_data(url)
        qr.make(fit=True)
        qr.make_image(fill_color="black", back_color="white").save(path)
        print(f"✓ {node_id:<6} {ntype:<12} {label[:42]:<42} -> {url}")

    print(f"\nDone. {len(nodes)} QR codes written to {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
